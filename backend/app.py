"""The HTTP surface. Routes and nothing else.

Everything is assembled in config.Services, so this file reads as a list of
what the API offers. Three things it does that are not just plumbing.

The translation sandwich: a question in another language is translated to
English on the way in, the answer is translated back on the way out, and the
English original is returned alongside it. The frontend should store that
English text as conversation history — it currently keeps the translated one
and sends it back, so on a French session the model reads its own French output
while being told never to mirror the language of the conversation.

The sources panel: an ECDC citation appears only when surveillance text
actually reached the model. It used to be built from the list of diseases the
scanner found, which is populated even when the data block comes back empty, so
a UK user was shown a citation to the ECDC Atlas for data that never existed.

Failures get status codes. A generation error used to be caught and returned as
the answer, at HTTP 200 — "Error: 429 rate limit exceeded" appeared in the chat
bubble while every monitor stayed green.
"""

import atexit
import io
import os
import time
from collections import defaultdict, deque
from contextlib import asynccontextmanager

from fastapi import FastAPI, File, Header, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse
from pydantic import BaseModel

import attribution
import config
import ecdc
import ontology
import logger as logging_
import pipeline
import prompts
import render
import report
import review
import translate

MAX_UPLOAD = 25 * 1024 * 1024


# --- request bodies ----------------------------------------------------------

class AskRequest(BaseModel):
    question: str
    country: str = "All Europe"
    language: str = "English"
    audience: str = "Policymaker"
    length: str = "Standard"
    session_id: str = ""
    history: list = []


class DocumentRequest(BaseModel):
    history: list
    audience: str = "Policymaker"
    country: str = "All Europe"
    language: str = "English"


class EvaluateRequest(BaseModel):
    text: str
    country: str = "All Europe"
    top_per_chunk: int = 3
    shortlist_n: int = 6
    min_score: float = 0.70


# --- optional gatekeeping ----------------------------------------------------

class RateLimit:
    """Requests per minute per client. Off unless RATE_LIMIT is set.

    In memory, so it resets when the process does and does not coordinate
    across replicas. That is honest for one small Space and better than the
    nothing that was there before.
    """

    def __init__(self, per_minute):
        self.per_minute = per_minute
        self.seen = defaultdict(deque)

    def check(self, client):
        if not self.per_minute:
            return
        now = time.time()
        recent = self.seen[client]
        while recent and now - recent[0] > 60:
            recent.popleft()
        if len(recent) >= self.per_minute:
            raise HTTPException(429, "Too many requests. Try again in a minute.")
        recent.append(now)


def create_app(services=None):
    holder = {"services": services}
    api_key = os.environ.get("API_KEY")
    limiter = RateLimit(config.env("RATE_LIMIT", 0, int))
    expose_debug = config.env("EXPOSE_DEBUG_ROUTES", False, config._boolean)

    def svc():
        """The assembled system.

        Built on first use rather than at import, so the module can be imported
        without a Mistral key — which is what lets the routes be tested at all.
        The holder rather than app.state, so nothing depends on the lifespan
        hook having run first.
        """
        if holder["services"] is None:
            holder["services"] = config.Services.build()
        return holder["services"]

    @asynccontextmanager
    async def lifespan(app):
        app.state.services = svc()          # build at startup, not on the first request
        # Belt and braces: lifespan shutdown runs on a clean stop, and a Space
        # being reclaimed is closer to a kill than to a clean stop.
        atexit.register(app.state.services.log.close)
        print(f"[APP] ready — {app.state.services.health()}")
        yield
        app.state.services.log.close()

    api = FastAPI(title="DSCAdapt API", version="1.0", lifespan=lifespan)
    api.add_middleware(
        CORSMiddleware,
        allow_origins=[o.strip() for o in
                       config.env("ALLOW_ORIGINS", "*").split(",")],
        allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

    def guard(request, key):
        if api_key and key != api_key:
            raise HTTPException(401, "Missing or invalid API key.")
        limiter.check(request.client.host if request.client else "unknown")

    # -- health --------------------------------------------------------------

    @api.head("/")
    @api.get("/")
    def health():
        """Reports what actually works, so a monitor can watch for degradation.

        The old health check said ok whenever the process was alive. If the
        Mistral key was revoked it still said ok while every answer failed.
        """
        state = svc().health()
        state["logger"] = svc().log.health()
        return state

    # -- chat ----------------------------------------------------------------

    @api.post("/ask")
    def ask(req: AskRequest, request: Request, x_api_key: str = Header(None)):
        guard(request, x_api_key)
        services = svc()
        target = translate.code_for(req.language)

        question = services.translator.to_english(req.question) if target else req.question

        try:
            answer = services.pipeline.respond(
                message=question, history=req.history, audience=req.audience,
                country=req.country, length=req.length)
        except pipeline.PipelineError as err:
            services.log.record(logging_.error_record(
                question, "pipeline", err, session=req.session_id,
                audience=req.audience, country=req.country,
                language=req.language, length=req.length))
            raise HTTPException(502, "The analysis could not be completed. Please try again.")

        # Attribution only makes sense in English: the source text is English,
        # so a translated answer would match nothing.
        marks = None
        display = answer.text
        if answer.route not in ("redirect", "empty"):
            if target:
                display = services.translator.answer(answer.text, target)
            else:
                marks = attribution.annotate(
                    answer.text, hits=answer.hits, web_text=answer.web_text,
                    surveillance_block=answer.ecdc_block, embed=services.embed)
                display = marks.html

        services.log.record(logging_.query_record(
            question, answer, audience=req.audience, country=req.country,
            language=req.language, length=req.length, session=req.session_id,
            turn=len([h for h in pipeline.clean_history(req.history)
                      if h["role"] == "user"]) + 1,
            attribution=marks,
            system_prompt=prompts.system_prompt(req.audience, req.country, req.length),
            capture_text=services.settings.log_text))

        return {
            "body": display,
            # Store this as history, not the translated body. See the note at
            # the top of this file.
            "english_answer": answer.text,
            "sources": build_sources(answer, req.country),
            "meta": {
                "words": len(display.split()),
                "passages": len(answer.hits),
                "web": "1 web search" if answer.web_text else "no web",
                "route": answer.route,
                "duration": f"{answer.timing.get('total', 0)}s",
                "rewritten": answer.rewritten,
            },
        }

    # -- documents -----------------------------------------------------------

    @api.post("/transcript")
    def transcript(req: DocumentRequest, request: Request, x_api_key: str = Header(None)):
        guard(request, x_api_key)
        try:
            path = render.render_transcript(req.history, req.audience, req.country, req.language)
        except Exception as err:
            print(f"[TRANSCRIPT] {err}")
            raise HTTPException(500, "The transcript could not be generated.")
        if not path:
            raise HTTPException(400, "No conversation to render.")
        return FileResponse(path, media_type="application/pdf",
                            filename=os.path.basename(path))

    @api.post("/briefing")
    def briefing(req: DocumentRequest, request: Request, x_api_key: str = Header(None)):
        guard(request, x_api_key)
        services = svc()
        try:
            result = render.Briefing(services.chat).build(req.history)
        except Exception as err:
            print(f"[BRIEFING] {err}")
            raise HTTPException(502, "The briefing could not be generated.")
        if result is None:
            raise HTTPException(400, "Not enough substantive conversation to generate a briefing.")

        sections = result.sections
        target = translate.code_for(req.language)
        if target:
            for field in ("summary", "cobenefits", "tradeoffs", "uncertainties"):
                setattr(sections, field,
                        services.translator.from_english(getattr(sections, field), target))

        # The draft matters more than the final text: the difference between
        # them is what the audit removed, which is the only way to tell whether
        # the audit is doing anything.
        entry = {"kind": "briefing", "turns": result.turns, "audited": result.audited,
                 "country": req.country, "language": req.language,
                 "questions": len([m for m in req.history if m.get("role") == "user"])}
        if services.settings.log_text:
            entry["draft"] = result.draft.as_text()
            entry["final"] = result.sections.as_text()
        services.log.record(entry)
        try:
            path = render.render_briefing(sections, req.history, req.audience,
                                          req.country, req.language)
        except Exception as err:
            print(f"[BRIEFING] render failed: {err}")
            raise HTTPException(500, "The briefing could not be rendered.")
        return FileResponse(path, media_type="application/pdf",
                            filename=os.path.basename(path))

    # -- reviewer ------------------------------------------------------------

    @api.post("/evaluate")
    def evaluate(req: EvaluateRequest, request: Request, x_api_key: str = Header(None)):
        guard(request, x_api_key)
        services = svc()
        try:
            result = services.reviewer.review(
                req.text, country=req.country, top_per_chunk=req.top_per_chunk,
                shortlist_n=req.shortlist_n, min_score=req.min_score)
        except ValueError as err:
            raise HTTPException(400, str(err))
        except Exception as err:
            services.log.record(logging_.error_record(req.text, "review", err,
                                                      country=req.country))
            raise HTTPException(502, "The review could not be completed.")

        # The record holds a length and a hash. Someone reviewing an
        # unpublished strategy should not find it in a dataset.
        services.log.record({"kind": "review",
                             **result.log_record(req.country, req.text,
                                                 include_comments=services.settings.log_text)})
        return {"note": result.note, "shortlist": result.shortlist,
                "comments": result.comments, "out_of_scope": result.out_of_scope}

    # -- extraction ----------------------------------------------------------

    @api.post("/extract")
    def extract(file: UploadFile = File(...)):
        name = (file.filename or "").lower()
        data = file.file.read()
        if len(data) > MAX_UPLOAD:
            return {"error": "File too large (max 25 MB). Paste the text instead."}
        try:
            if name.endswith(".pdf"):
                text = extract_pdf(data)
            elif name.endswith(".docx"):
                text = extract_docx(data)
            elif name.endswith(".txt"):
                text = data.decode("utf-8", "ignore").strip()
            else:
                return {"error": "Unsupported file type. Upload a PDF, DOCX, or TXT."}
        except Exception as err:
            return {"error": "Could not read this file. It may be scanned, corrupted, or "
                             f"password-protected. ({type(err).__name__})"}
        if not text:
            return {"error": "No readable text found. A scanned PDF needs OCR, which is not "
                             "supported here. Paste the text instead."}

        result = {"text": text, "chars": len(text), "filename": file.filename}
        # Said here rather than after the review call, so the length problem is
        # visible before anything is submitted.
        if len(text) > review.MAX_CHARS:
            result["warning"] = (
                f"This is {len(text):,} characters and the reviewer takes about "
                f"{review.MAX_CHARS:,}. Paste the relevant section instead.")
        return result

    # -- traces ---------------------------------------------------------------

    @api.get("/logs", response_class=HTMLResponse)
    def logs(limit: int = 200, x_api_key: str = Header(None)):
        """The readable version of the traces.

        Rendered when somebody opens it, rather than rebuilt inside every
        request the way the original was — which is why the original could not
        be asked to record more than it did.
        """
        if api_key and x_api_key != api_key:
            raise HTTPException(401, "Missing or invalid API key.")
        return report.render(report.read(svc().settings.logs_dir, limit=limit))

    # -- debug ---------------------------------------------------------------

    if expose_debug:
        @api.get("/debug/graph")
        def debug_graph(country: str = "All Europe"):
            """Which diseases each measure reaches for a country, and why."""
            services = svc()
            reachable = {}
            for measure in ontology.MEASURE_RISKS:
                diseases = [d for d in ontology.diseases_for(measure)
                            if services.surveillance.relevant(d, country)]
                if diseases:
                    reachable[measure] = diseases
            return {"country": country,
                    "surveillance": services.surveillance.available,
                    "covered": services.surveillance.has_country(country),
                    "years": services.surveillance.year_range,
                    "reachable": reachable}

        @api.get("/debug/disease")
        def debug_disease(disease: str):
            """Cumulative reported cases per country for one disease.

            How the autochthonous and endemic country sets get checked against
            what the Atlas actually holds. Dropped in the rewrite; it is the
            only tool for calibrating those lists.
            """
            surveillance = svc().surveillance
            if not surveillance.available:
                return {"error": "surveillance data not loaded"}
            record = ecdc.DISEASES.get(disease, {})
            totals = surveillance.country_totals(disease)
            if totals is None:
                return {"disease": disease, "note": "not in the extract",
                        "known": surveillance.known_diseases()}
            return {"disease": disease, "resolved_to": record.get("ecdc", disease),
                    "years": surveillance.years, "countries": totals,
                    "gated_to": record.get("countries")}

        @api.get("/debug/corpus")
        def debug_corpus():
            services = svc()
            return {"chunks": len(services.index),
                    "report": services.corpus.report().split("\n")}

    return api


# --- helpers -----------------------------------------------------------------

def build_sources(answer, country):
    """What the sources panel shows.

    The ECDC row is keyed on whether a data block was built, not on whether a
    disease name was spotted. Those come apart for any country the Atlas does
    not cover, and the panel then cited data that was never used.
    """
    sources = [{"type": "evidence", "score": round(hit.score, 2),
                "title": hit.chunk.section,
                "origin": "Evidence base · " + hit.chunk.source
                          .replace(".docx.txt", "").replace(".txt", "")}
               for hit in answer.hits]

    if answer.web_text:
        sources.append({"type": "web", "score": None,
                        "title": "Current policy and geographic context",
                        "origin": "Web search"})

    if answer.used_surveillance:
        sources.append({"type": "ecdc", "score": None,
                        "title": f"Surveillance data: {', '.join(answer.diseases[:5])}",
                        "origin": f"ECDC Surveillance Atlas · {country}"})

    for record, score in answer.scoping or []:
        sources.append({"type": "evidence", "score": round(score, 2),
                        "title": record["section"],
                        "origin": "IDAlert D2.2 scoping review (2024)"})
    return sources


def extract_pdf(data):
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = [(page.extract_text() or "").strip() for page in pdf.pages]
    return "\n\n".join(p for p in pages if p).strip()


def extract_docx(data):
    import docx
    document = docx.Document(io.BytesIO(data))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts).strip()


api = create_app()


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(api, host="0.0.0.0", port=7860)
